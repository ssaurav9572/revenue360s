from __future__ import annotations

import csv
import io
import json
import re
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from auth import user_from_header
from db import LEAD_STAGES, execute, format_tags, one, query
import learning
from engine import enroll, find_or_create_lead

router = APIRouter(prefix="/api/leads", tags=["leads"])


@router.get("")
def list_leads(user: dict = Depends(user_from_header)) -> list:
    rows = query("SELECT * FROM leads WHERE workspace_id = ? ORDER BY created_at DESC LIMIT 300", (user["workspace_id"],))
    for row in rows:
        row["tags"] = format_tags(row.get("tags"))
        try:
            row["custom_fields"] = json.loads(row.get("custom_fields") or "{}")
        except (TypeError, json.JSONDecodeError):
            row["custom_fields"] = {}
    return rows


@router.post("")
def add_lead(payload: dict, user: dict = Depends(user_from_header)) -> dict:
    ws = user["workspace_id"]
    if not payload.get("phone") and not payload.get("email"):
        raise HTTPException(400, "Phone or email is required.")
    tags = [t.strip() for t in str(payload.get("tags") or "").split(",") if t.strip()]
    lead, created = find_or_create_lead(ws, payload.get("name", ""), payload.get("phone", ""), payload.get("email", ""), payload.get("source", "manual"), tags)
    if payload.get("sequence_id") and lead.get("id"):
        enroll(ws, payload["sequence_id"], lead["id"])
    return {"lead": lead, "created": created}


@router.post("/import/preview")
async def preview_leads(file: UploadFile = File(...), user: dict = Depends(user_from_header)) -> dict:
    del user
    filename = (file.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx")):
        raise HTTPException(400, "Choose a CSV or XLSX file.")
    raw = await file.read()
    if len(raw) > 8_000_000:
        raise HTTPException(400, "File too large (8MB max).")

    workbook = None
    try:
        if filename.endswith(".csv"):
            text = raw.decode("utf-8-sig")
            source_rows = csv.reader(io.StringIO(text))
        else:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            source_rows = sheet.iter_rows(values_only=True)
        values = []
        for row in source_rows:
            row = list(row)
            if not any(value is not None and str(value).strip() for value in row):
                continue
            if len(row) > 100:
                raise HTTPException(400, "Imports are limited to 100 columns.")
            if values and len(values) > 5000:
                raise HTTPException(400, "Imports are limited to 5,000 rows.")
            values.append(row)
    except (UnicodeDecodeError, ValueError, csv.Error, BadZipFile, InvalidFileException) as error:
        raise HTTPException(400, "Could not read this file. Check its format and try again.") from error
    finally:
        if workbook:
            workbook.close()

    values = [row for row in values if any(value is not None and str(value).strip() for value in row)]
    if not values:
        raise HTTPException(400, "The file has no rows to import.")
    if len(values) - 1 > 5000 or len(values[0]) > 100:
        raise HTTPException(400, "Imports are limited to 5,000 rows and 100 columns.")

    def cell_text(value: object) -> str:
        if value is None:
            return ""
        if hasattr(value, "isoformat"):
            return value.isoformat()
        return str(value)

    columns = [cell_text(value).strip() or f"Column {index + 1}" for index, value in enumerate(values[0])]
    rows = [[cell_text(value) for value in row[:len(columns)]] + [""] * max(0, len(columns) - len(row)) for row in values[1:]]
    return {"columns": columns, "rows": rows}


@router.post("/import")
def import_leads(payload: dict, user: dict = Depends(user_from_header)) -> dict:
    columns = payload.get("columns")
    rows = payload.get("rows")
    if not isinstance(columns, list) or not columns or len(columns) > 100 or not isinstance(rows, list) or len(rows) > 5000:
        raise HTTPException(400, "Provide up to 5,000 rows and 100 columns.")

    aliases = {
        "name": {"name", "fullname", "contactname", "leadname"},
        "phone": {"phone", "mobile", "whatsapp", "phonenumber", "telephone"},
        "email": {"email", "emailaddress"},
        "source": {"source", "leadsource"},
        "tags": {"tags", "tag"},
        "notes": {"notes", "note"},
    }
    canonical = {}
    unique_columns = []
    used = set()
    for index, column in enumerate(columns):
        label = str(column or "").strip() or f"Column {index + 1}"
        key = label
        suffix = 2
        while key.casefold() in used:
            key = f"{label} ({suffix})"
            suffix += 1
        used.add(key.casefold())
        unique_columns.append(key)
        normalized = re.sub(r"[^a-z0-9]", "", label.casefold())
        for target, names in aliases.items():
            if normalized in names:
                canonical[target] = index
                break

    created = updated = skipped = 0
    errors = []
    for row_number, row in enumerate(rows, start=2):
        if not isinstance(row, list) or len(row) > len(columns):
            errors.append({"row": row_number, "message": "Row has an invalid shape."})
            continue
        values = [str(value if value is not None else "").strip() for value in row]
        if not any(values):
            skipped += 1
            continue
        values.extend([""] * (len(columns) - len(values)))
        get_value = lambda field: values[canonical[field]] if field in canonical else ""
        name, phone, email = get_value("name"), get_value("phone"), get_value("email")
        if not phone and not email:
            errors.append({"row": row_number, "message": "Phone or email is required."})
            continue
        if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            errors.append({"row": row_number, "message": "Email address is invalid."})
            continue
        if len(name) > 255 or len(phone) > 64 or len(email) > 255 or len(get_value("source")) > 80:
            errors.append({"row": row_number, "message": "Name, phone, or email exceeds the allowed length."})
            continue

        custom_fields = {unique_columns[index]: value for index, value in enumerate(values) if value and index not in canonical.values()}
        tags = [tag.strip() for tag in get_value("tags").split(",") if tag.strip()]
        workspace_id = user["workspace_id"]
        lead, is_new = find_or_create_lead(
            workspace_id, name, phone, email, get_value("source") or "import", tags
        )
        if custom_fields or get_value("notes"):
            try:
                existing_custom = json.loads(lead.get("custom_fields") or "{}")
            except (TypeError, json.JSONDecodeError):
                existing_custom = {}
            existing_custom.update(custom_fields)
            execute(
                "UPDATE leads SET custom_fields = ?, notes = CASE WHEN ? <> '' THEN ? ELSE notes END WHERE id = ? AND workspace_id = ?",
                (json.dumps(existing_custom), get_value("notes"), get_value("notes"), lead["id"], workspace_id),
            )
        created += int(is_new)
        updated += int(not is_new)
    return {"created": created, "updated": updated, "skipped": skipped, "errors": errors}


@router.post("/{lead_id}/stage")
def set_stage(lead_id: str, payload: dict, user: dict = Depends(user_from_header)) -> dict:
    ws = user["workspace_id"]
    stage = str(payload.get("stage") or "new")
    if stage not in LEAD_STAGES:
        raise HTTPException(400, "Unknown stage.")
    lead = one("SELECT stage FROM leads WHERE id = ? AND workspace_id = ?", (lead_id, ws))
    if not lead:
        raise HTTPException(404, "Lead not found.")
    execute("UPDATE leads SET stage = ?, status = ? WHERE id = ? AND workspace_id = ?", (stage, stage, lead_id, ws))
    if lead["stage"] != stage:
        try:
            learning.record_stage_feedback(ws, lead_id, stage)
        except Exception:
            pass
    return {"ok": True}


@router.post("/{lead_id}/notes")
def set_notes(lead_id: str, payload: dict, user: dict = Depends(user_from_header)) -> dict:
    execute("UPDATE leads SET notes = ? WHERE id = ? AND workspace_id = ?", (payload.get("notes") or "", lead_id, user["workspace_id"]))
    return {"ok": True}


@router.get("/{lead_id}/timeline")
def timeline(lead_id: str, user: dict = Depends(user_from_header)) -> dict:
    lead = one("SELECT * FROM leads WHERE id = ? AND workspace_id = ?", (lead_id, user["workspace_id"]))
    if not lead:
        raise HTTPException(404, "Lead not found.")
    try:
        lead["custom_fields"] = json.loads(lead.get("custom_fields") or "{}")
    except (TypeError, json.JSONDecodeError):
        lead["custom_fields"] = {}
    messages = query(
        """
        SELECT created_at, channel, direction, subject, body, status
        FROM messages
        WHERE workspace_id = ? AND (lead_id = ? OR recipient IN (?, ?))
        ORDER BY created_at DESC LIMIT 50
        """,
        (user["workspace_id"], lead_id, lead.get("phone") or "", lead.get("email") or ""),
    )
    return {"lead": lead, "messages": messages, "stages": list(LEAD_STAGES)}